from openpyxl import load_workbook

from app.core.convert_from_pdf import pdf_to_excel, pdf_to_jpg, pdf_to_word


def test_pdf_to_jpg(make_pdf, tmp_path):
    source = make_pdf("source.pdf", pages=3)
    output_dir = tmp_path / "pages"

    outputs = pdf_to_jpg(source, output_dir)

    assert len(outputs) == 3
    for out_path in outputs:
        assert out_path.exists()


def test_pdf_to_word(make_pdf, tmp_path):
    source = make_pdf("source.pdf", pages=1, text="Hello PDFLocal")
    output = tmp_path / "converted.docx"

    pdf_to_word(source, output)

    assert output.exists()
    assert output.stat().st_size > 0


def test_pdf_to_excel(make_pdf, tmp_path):
    source = make_pdf("source.pdf", pages=2, text="Some text content")
    output = tmp_path / "converted.xlsx"

    pdf_to_excel(source, output)

    workbook = load_workbook(output)
    assert len(workbook.sheetnames) == 2
