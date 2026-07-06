from pathlib import Path

import fitz
import pdfplumber
from openpyxl import Workbook


def pdf_to_jpg(path: Path, output_dir: Path, dpi: int = 150) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = []
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    with fitz.open(path) as doc:
        for index, page in enumerate(doc):
            pixmap = page.get_pixmap(matrix=matrix)
            out_path = output_dir / f"page_{index + 1}.jpg"
            pixmap.save(out_path)
            outputs.append(out_path)
    return outputs


def pdf_to_word(path: Path, output: Path) -> Path:
    from pdf2docx import Converter

    converter = Converter(str(path))
    try:
        converter.convert(str(output))
    finally:
        converter.close()
    return output


def pdf_to_excel(path: Path, output: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages):
            sheet = workbook.create_sheet(title=f"Page{index + 1}")
            tables = page.extract_tables()
            if tables:
                for table in tables:
                    for row in table:
                        sheet.append(row)
            else:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    sheet.append([line])
    if not workbook.sheetnames:
        workbook.create_sheet(title="Sheet1")
    workbook.save(output)
    return output
